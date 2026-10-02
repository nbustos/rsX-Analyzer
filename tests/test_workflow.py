"""Focused tests for the compact workflow orchestration API and PDF export.

These use the bundled ``atlas-4S356Parcels_dseg.tsv`` fixture (a real 4S atlas
sidecar shipped with the repository) plus synthetic, randomly generated
connectivity matrices and motion HDF5 files, since no real sample dataset is
checked into the repository. The synthetic fixtures are built to reproduce
the same *shape* of outputs the notebook reports against its external sample
data (403 connectivity features for the bundled 356-parcel 4S atlas, matched
subject/session/run averaging, and an explicit unreadable-motion-file
inventory).
"""
import re
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd

from rsx_analyzer import (
    average_matched_runs,
    derive_analysis_atlas,
    run_workflow,
    robust_outlier_screen,
    triple_network_summary,
)

try:
    import h5py
    HAVE_H5PY = True
except ImportError:
    HAVE_H5PY = False

try:
    import matplotlib  # noqa: F401
    HAVE_MATPLOTLIB = True
except ImportError:
    HAVE_MATPLOTLIB = False

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

ATLAS_PATH = Path(__file__).parents[1] / "atlases" / "atlas-4S356Parcels_dseg.tsv"


def _write_conn_matrix(directory, atlas_names, subject, session, task, run, rng):
    n = len(atlas_names)
    values = rng.normal(size=(n, n))
    matrix = (values + values.T) / 2
    np.fill_diagonal(matrix, 1.0)
    frame = pd.DataFrame(matrix, index=atlas_names, columns=atlas_names)
    run_entity = f"_run-{run}" if run is not None else ""
    name = f"sub-{subject}_ses-{session}_task-{task}{run_entity}_conmat.tsv"
    frame.to_csv(Path(directory) / name, sep="\t")


def _write_motion_hdf5(directory, subject, session, task, run, remaining_seconds_at_03):
    run_entity = f"_run-{run}" if run is not None else ""
    name = f"sub-{subject}_ses-{session}_task-{task}{run_entity}_motion.hdf5"
    with h5py.File(Path(directory) / name, "w") as handle:
        for threshold in (0.3, 0.4, 0.5):
            group = handle.create_group(f"dcan_motion/fd_{threshold:g}")
            group.create_dataset("threshold", data=threshold)
            group.create_dataset("total_frame_count", data=100.0)
            group.create_dataset("remaining_total_frame_count", data=90.0 - threshold)
            group.create_dataset("remaining_frame_mean_FD", data=0.1)
            group.create_dataset(
                "remaining_seconds", data=remaining_seconds_at_03 - threshold * 10
            )
            group.create_dataset("skip", data=0.0)


def _write_linc_motion_csv(path, rows):
    pd.DataFrame(rows, columns=[
        "sub", "ses", "task", "run", "space", "res", "desc",
        "mean_fd", "num_censored_volumes", "num_retained_volumes",
    ]).to_csv(path, index=False)


@unittest.skipUnless(HAVE_H5PY, "requires optional dependency h5py")
class RunWorkflowTests(unittest.TestCase):
    def test_reproduces_feature_count_and_matched_run_rows(self):
        atlas = derive_analysis_atlas(ATLAS_PATH)
        rng = np.random.default_rng(0)
        with tempfile.TemporaryDirectory() as directory:
            conn_dir = Path(directory) / "conn_mats"
            motion_dir = Path(directory) / "dcan_qc"
            conn_dir.mkdir()
            motion_dir.mkdir()
            subjects = [f"{i:02d}" for i in range(1, 6)]
            for offset, subject in enumerate(subjects):
                _write_conn_matrix(conn_dir, atlas.names, subject, "A", "rest", "1", rng)
                _write_motion_hdf5(motion_dir, subject, "A", "rest", "1", 277.0 - offset)
            # An unreadable motion file must remain visible, not silently dropped.
            (motion_dir / "sub-99_ses-A_task-rest_run-1_motion.hdf5").write_bytes(b"not hdf5")

            result = run_workflow(conn_dir, motion_dir, ATLAS_PATH)

        self.assertEqual(len(result.connectivity.feature_columns), 403)
        self.assertEqual(len(result.discovery.conn_files), 5)
        self.assertEqual(len(result.discovery.motion_files), 6)
        self.assertEqual(len(result.motion.motion_df), 5)
        self.assertEqual(len(result.motion.motion_load_errors), 1)
        self.assertIn("99", result.motion.motion_load_errors["subject"].tolist())
        self.assertEqual(len(result.master.master_df), 5)
        self.assertEqual(len(result.merged_runs), 5)
        self.assertFalse(result.run_level.missingness_report.empty)
        self.assertIn("n_subjects_with_both_readable", result.run_level.coverage_summary.columns)
        self.assertEqual(set(result.threshold_table["fd_threshold"]), {0.3, 0.4, 0.5})
        self.assertIn("SubjectID", result.master.master_df.columns)
        self.assertIn("remaining_seconds", result.master.master_df.columns)

    def test_average_matched_runs_merges_repeated_runs_and_screens_outliers(self):
        atlas = derive_analysis_atlas(ATLAS_PATH)
        rng = np.random.default_rng(1)
        with tempfile.TemporaryDirectory() as directory:
            conn_dir = Path(directory) / "conn_mats"
            motion_dir = Path(directory) / "dcan_qc"
            conn_dir.mkdir()
            motion_dir.mkdir()
            layout = {"01": ["1", "2"], "02": ["1"], "03": ["1"]}
            for offset, (subject, runs) in enumerate(layout.items()):
                for run in runs:
                    _write_conn_matrix(conn_dir, atlas.names, subject, "A", "rest", run, rng)
                    _write_motion_hdf5(motion_dir, subject, "A", "rest", run, 277.0 - offset - int(run))

            result = run_workflow(conn_dir, motion_dir, ATLAS_PATH)

        self.assertEqual(len(result.master.master_df), 4)
        self.assertEqual(len(result.merged_runs), 3)
        subject_01 = result.merged_runs.loc[result.merged_runs["SubjectID"] == "01"].iloc[0]
        self.assertEqual(subject_01["n_runs_included"], 2)
        self.assertEqual(subject_01["runs_included"], "01, 02")

        screen = robust_outlier_screen(
            result.merged_runs, result.master.feature_columns,
            subject_column="SubjectID", session_column="Session",
        )
        self.assertEqual(
            set(screen), {"scores", "mask", "feature_summary", "feature_flags",
                          "subject_report", "unscorable"},
        )
        self.assertEqual(len(screen["feature_summary"]), len(result.master.feature_columns))

    def test_averages_two_longest_runs_and_keeps_runless_single_run_session(self):
        atlas = derive_analysis_atlas(ATLAS_PATH)
        rng = np.random.default_rng(3)
        with tempfile.TemporaryDirectory() as directory:
            conn_dir = Path(directory) / "conn_mats"
            motion_dir = Path(directory) / "dcan_qc"
            conn_dir.mkdir()
            motion_dir.mkdir()
            for run, seconds in (("1", 200.0), ("2", 400.0), ("3", 300.0)):
                _write_conn_matrix(conn_dir, atlas.names, "01", "A", "rest", run, rng)
                _write_motion_hdf5(motion_dir, "01", "A", "rest", run, seconds + 3.0)
            _write_conn_matrix(conn_dir, atlas.names, "02", "A", "rest", None, rng)
            _write_motion_hdf5(motion_dir, "02", "A", "rest", None, 300.0)

            result = run_workflow(conn_dir, motion_dir, ATLAS_PATH)

        self.assertEqual(len(result.master.master_df), 4)
        subject_01_master = result.master.master_df.loc[
            result.master.master_df["SubjectID"] == "01"
        ]
        selected = result.merged_runs.loc[
            result.merged_runs["SubjectID"] == "01"
        ].iloc[0]
        self.assertEqual(selected["n_runs_included"], 2)
        self.assertEqual(selected["runs_included"], "02, 03")
        self.assertEqual(selected["remaining_seconds"], 350.0)
        feature = result.master.feature_columns[0]
        expected_feature_mean = subject_01_master.loc[
            subject_01_master["Run"].isin(["02", "03"]), feature
        ].mean()
        self.assertAlmostEqual(selected[feature], expected_feature_mean)

        single_run = result.merged_runs.loc[
            result.merged_runs["SubjectID"] == "02"
        ].iloc[0]
        self.assertEqual(single_run["n_runs_included"], 1)
        self.assertEqual(single_run["runs_included"], "n/a")
        self.assertEqual(single_run["remaining_seconds"], 297.0)

    def test_average_matched_runs_rejects_nonpositive_run_limit(self):
        frame = pd.DataFrame({
            "SubjectID": ["01"], "Session": ["A"], "Run": ["1"],
            "remaining_seconds": [100.0], "feature": [0.5],
        })
        with self.assertRaisesRegex(ValueError, "max_runs"):
            average_matched_runs(frame, ["feature"], motion_feature_columns=(), max_runs=0)

    def test_average_matched_runs_requires_all_selected_runs_to_meet_threshold(self):
        frame = pd.DataFrame({
            "SubjectID": ["01", "01", "01", "02", "02", "03"],
            "Session": ["A"] * 6,
            "Run": ["01", "02", "03", "01", "02", "n/a"],
            "remaining_seconds": [240, 260, 500, 300, 239, 240],
            "feature": [1.0, 3.0, 100.0, 5.0, 9.0, 7.0],
        })
        result = average_matched_runs(
            frame, ["feature"], motion_feature_columns=(),
            minimum_selection_value=240,
        )
        self.assertEqual(result["SubjectID"].tolist(), ["01", "03"])
        self.assertEqual(result.loc[0, "runs_included"], "02, 03")
        self.assertEqual(result.loc[0, "feature"], 51.5)
        self.assertEqual(result.loc[1, "runs_included"], "n/a")
        self.assertEqual(result.loc[1, "n_runs_included"], 1)

    @unittest.skipUnless(HAVE_H5PY, "requires optional dependency h5py")
    def test_linc_csv_mode_loads_metadata_and_selects_most_retained_runs(self):
        atlas = derive_analysis_atlas(ATLAS_PATH)
        rng = np.random.default_rng(4)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            conn_dir = root / "conn_mats"
            motion_dir = root / "linc_qc"
            conn_dir.mkdir()
            motion_dir.mkdir()
            for run in ("1", "2", "3"):
                _write_conn_matrix(conn_dir, atlas.names, "01", "A", "rest", run, rng)
            _write_conn_matrix(conn_dir, atlas.names, "02", "A", "rest", None, rng)
            _write_linc_motion_csv(motion_dir / "linc_metrics.csv", [
                ("sub-01", "ses-A", "task-rest", "run-1", "MNI", "2", "dcan", "0.20", "10", "350"),
                ("sub-01", "ses-A", "task-rest", "run-2", "MNI", "2", "dcan", "0.10", "5", "400"),
                ("sub-01", "ses-A", "task-rest", "run-3", "MNI", "2", "dcan", "0.15", "8", "300"),
                ("sub-02", "ses-A", "task-rest", None, "MNI", "2", "dcan", "0.12", "7", "250"),
            ])

            result = run_workflow(conn_dir, motion_dir, ATLAS_PATH)

        self.assertEqual(result.motion_mode, "linc_qc")
        self.assertEqual(result.run_selection_column, "num_retained_volumes")
        self.assertEqual(len(result.motion.motion_df), 4)
        self.assertEqual(len(result.master.master_df), 4)
        self.assertIn("mean_fd", result.master.master_df)
        self.assertIn("num_censored_volumes", result.master.master_df)
        self.assertIn("num_retained_volumes", result.master.master_df)
        self.assertEqual(
            result.master.master_df.loc[
                result.master.master_df["SubjectID"] == "01", "Run"
            ].tolist(),
            ["01", "02", "03"],
        )
        subject_01 = result.merged_runs.loc[
            result.merged_runs["SubjectID"] == "01"
        ].iloc[0]
        self.assertEqual(subject_01["runs_included"], "01, 02")
        self.assertEqual(subject_01["num_retained_volumes"], 750.0)
        self.assertEqual(subject_01["num_censored_volumes"], 15.0)
        self.assertAlmostEqual(subject_01["mean_fd"], 0.15)
        self.assertEqual(subject_01["space"], "MNI")
        self.assertEqual(subject_01["res"], "2")
        self.assertEqual(subject_01["desc"], "dcan")
        single_run = result.merged_runs.loc[
            result.merged_runs["SubjectID"] == "02"
        ].iloc[0]
        self.assertEqual(single_run["runs_included"], "n/a")
        self.assertEqual(single_run["num_retained_volumes"], 250.0)
        self.assertTrue((result.run_level.run_level_df["file_status"] == "both files readable").all())

    def test_motion_mode_is_selected_from_motion_folder_name(self):
        from rsx_analyzer import detect_motion_mode

        self.assertEqual(detect_motion_mode("/inputs/dcan_qc"), "dcan_qc")
        self.assertEqual(detect_motion_mode("/inputs/LINC_QC"), "linc_qc")
        with self.assertRaisesRegex(ValueError, "dcan_qc.*linc_qc"):
            detect_motion_mode("/inputs/motion")


@unittest.skipUnless(HAVE_H5PY and HAVE_MATPLOTLIB, "requires h5py and matplotlib")
class ExportAnalysisResultsPdfTests(unittest.TestCase):
    def test_produces_seven_page_print_ready_report(self):
        from rsx_analyzer import export_analysis_results_pdf

        atlas = derive_analysis_atlas(ATLAS_PATH)
        rng = np.random.default_rng(2)
        with tempfile.TemporaryDirectory() as directory:
            conn_dir = Path(directory) / "conn_mats"
            motion_dir = Path(directory) / "dcan_qc"
            conn_dir.mkdir()
            motion_dir.mkdir()
            for offset, subject in enumerate(["01", "02", "03"]):
                _write_conn_matrix(conn_dir, atlas.names, subject, "A", "rest", "1", rng)
                _write_motion_hdf5(motion_dir, subject, "A", "rest", "1", 277.0 - offset)

            result = run_workflow(conn_dir, motion_dir, ATLAS_PATH)
            screen = robust_outlier_screen(
                result.merged_runs, result.master.feature_columns,
                subject_column="SubjectID", session_column="Session",
            )
            network19_counts = pd.Series(
                atlas.networks["network_label_19network"]
            ).value_counts().to_dict()
            systems = {
                "Control (CEN)": {"network9": ("Cont",), "network19": ("ContA", "ContB", "ContC")},
                "Default (DMN)": {"network9": ("Default",), "network19": ("DefaultA", "DefaultB", "DefaultC")},
                "Salience (SN)": {"network9": ("SalVentAttn",), "network19": ("SalVentAttnA", "SalVentAttnB")},
            }
            long_values, triple_summary = triple_network_summary(
                result.merged_runs, network19_counts, systems
            )
            order = [f"{label} within" for label in systems] + [
                f"{a} x {b}" for i, a in enumerate(systems) for b in list(systems)[i + 1:]
            ]

            pdf_path = Path(directory) / "rs-X1_results.pdf"
            export_analysis_results_pdf(
                pdf_path,
                atlas_name=ATLAS_PATH.stem,
                n_atlas_parcels=len(atlas.names),
                fd_threshold=0.3,
                n_conn_files=len(result.discovery.conn_files),
                n_motion_files=len(result.discovery.motion_files),
                n_motion_readable=len(result.motion.motion_df),
                n_master_rows=len(result.master.master_df),
                n_merged_run_rows=len(result.merged_runs),
                n_features=len(result.master.feature_columns),
                n_flagged_subjects=len(screen["subject_report"]),
                threshold_summary=result.threshold_table,
                motion_associations=result.motion_associations,
                outlier_feature_summary=screen["feature_summary"],
                outlier_subject_report=screen["subject_report"],
                triple_summary=triple_summary,
                triple_network_df=long_values,
                triple_feature_order=order,
            )

            self.assertTrue(pdf_path.exists())
            page_markers = re.findall(rb"/Type\s*/Page[^s]", pdf_path.read_bytes())
            self.assertEqual(len(page_markers), 7)

    def test_requires_triple_network_figure_or_data(self):
        from rsx_analyzer import export_analysis_results_pdf

        empty = pd.DataFrame({"feature": [], "spearman_rho": []})
        with tempfile.TemporaryDirectory() as directory, self.assertRaises(ValueError):
            export_analysis_results_pdf(
                Path(directory) / "out.pdf",
                atlas_name="atlas", n_atlas_parcels=1, fd_threshold=0.3,
                n_conn_files=0, n_motion_files=0, n_motion_readable=0,
                n_master_rows=0, n_merged_run_rows=0, n_features=0,
                n_flagged_subjects=0, threshold_summary=empty,
                motion_associations=empty, outlier_feature_summary=empty,
                outlier_subject_report=empty, triple_summary=empty,
            )


@unittest.skipUnless(HAVE_MATPLOTLIB, "requires optional dependency matplotlib")
@unittest.skipUnless(HAVE_OPENPYXL, "requires optional dependency openpyxl")
class ExportExcelFigureTests(unittest.TestCase):
    def test_embeds_triple_network_plot_with_header_styling(self):
        import matplotlib.pyplot as plt
        from openpyxl import load_workbook
        from rsx_analyzer import export_excel

        master = pd.DataFrame({"SubjectID": ["01", "01"], "Run": ["1", "2"]})
        merged_runs = pd.DataFrame({"SubjectID": ["01", "02"], "value": [1.0, 2.0]})
        qc = pd.DataFrame({"SubjectID": ["01"], "n_flagged_values": [3]})
        triple_summary = pd.DataFrame({
            "partition": ["9-network"], "feature": ["Control within"], "mean": [0.5],
        })
        figure, ax = plt.subplots()
        ax.plot([0, 1], [0, 1])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rs-X1_analysis.xlsx"
            export_excel(
                {"master": master, "merged_runs": merged_runs, "QC": qc, "Trinetwork": triple_summary},
                path,
                figure=figure,
                figure_sheet="Trinetwork",
                figure_title="Triple-network connectivity by atlas resolution",
                figure_subtitle="19-network summaries use parcel-pair-weighted means.",
            )
            plt.close(figure)

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["master", "merged_runs", "QC", "Trinetwork"])
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                self.assertEqual(worksheet.freeze_panes, "A2")
                self.assertTrue(worksheet.auto_filter.ref)
                self.assertTrue(worksheet["A1"].font.bold)
            self.assertEqual(len(workbook["Trinetwork"]._images), 1)
            self.assertEqual(workbook["master"].max_row, 3)
            self.assertEqual(len(workbook["merged_runs"]._images), 0)

    def test_without_figure_still_applies_header_styling(self):
        from openpyxl import load_workbook
        from rsx_analyzer import export_excel

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "out.xlsx"
            export_excel({"merged_runs": pd.DataFrame({"a": [1]})}, path)
            worksheet = load_workbook(path)["merged_runs"]
            self.assertEqual(worksheet.freeze_panes, "A2")
            self.assertEqual(len(worksheet._images), 0)


@unittest.skipUnless(HAVE_MATPLOTLIB, "requires optional dependency matplotlib")
class PlotOutlierSummaryTests(unittest.TestCase):
    def test_renders_feature_and_subject_bar_charts(self):
        from rsx_analyzer import plot_outlier_summary

        feature_summary = pd.DataFrame({
            "feature": ["A_within_9net", "B_within_9net", "C_within_9net"],
            "n_outlier_rows": [5, 0, 2],
        })
        subject_report = pd.DataFrame({
            "SubjectID": ["01", "02", "03"],
            "n_flagged_values": [4, 7, 1],
        })

        feature_ax, subject_ax = plot_outlier_summary(feature_summary, subject_report, top_n=20)

        self.assertEqual(feature_ax.get_title(), "Features with the most robust outlier flags")
        self.assertEqual(subject_ax.get_title(), "Subjects with the most connectivity-feature outliers")
        # The zero-flag feature must be excluded from the feature panel.
        feature_labels = [label.get_text() for label in feature_ax.get_yticklabels()]
        self.assertNotIn("B_within_9net", feature_labels)
        subject_labels = [label.get_text() for label in subject_ax.get_yticklabels()]
        self.assertEqual(set(subject_labels), {"01", "02", "03"})

    def test_accepts_supplied_axes_pair(self):
        import matplotlib.pyplot as plt
        from rsx_analyzer import plot_outlier_summary

        feature_summary = pd.DataFrame({"feature": ["A"], "n_outlier_rows": [1]})
        subject_report = pd.DataFrame({"SubjectID": ["01"], "n_flagged_values": [2]})
        _, axes = plt.subplots(1, 2)

        feature_ax, subject_ax = plot_outlier_summary(
            feature_summary, subject_report, axes=axes
        )

        self.assertIs(feature_ax, axes[0])
        self.assertIs(subject_ax, axes[1])


if __name__ == "__main__":
    unittest.main()
