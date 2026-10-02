"""Optional plotting and report export utilities.

All print-ready PDF report generation lives here: small table/figure page
helpers plus :func:`export_analysis_results_pdf`, a high-level exporter that
reproduces the notebook's 7-page ``rs-X1_results.pdf`` report from computed
dataframes (and either a supplied triple-network figure or the long-form
triple-network values needed to build one internally).
"""
from __future__ import annotations
from collections.abc import Sequence
from pathlib import Path
import pandas as pd

DEFAULT_TRIPLE_PALETTE = {"9-network": "#315f7d", "19-network": "#d08b4f"}


def add_report_footer(fig: object, page_number: int, label: str = "rs-X1 analysis") -> None:
    """Add a consistent footer to a matplotlib figure."""
    fig.text(0.5, 0.018, f"{label} | Page {page_number}", ha="center", fontsize=8, color="#687782")


def add_report_table_page(pdf: object, title: str, dataframe: pd.DataFrame, page_number: int,
                          subtitle: str | None = None, max_column_chars: int = 72,
                          report_label: str = "rs-X1 analysis") -> None:
    """Render a compact dataframe page into a matplotlib PdfPages-like writer."""
    try:
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise ImportError("add_report_table_page requires optional dependency 'matplotlib'") from exc
    fig, ax = plt.subplots(figsize=(11.7, 8.3)); ax.axis("off")
    ax.set_title(title, loc="left", fontsize=18, weight="bold", color="#24465c", pad=24)
    if subtitle: fig.text(0.06, 0.91, subtitle, ha="left", fontsize=9, color="#566975")
    table_data = dataframe.copy()
    for col in table_data.select_dtypes(include="number"):
        table_data[col] = table_data[col].map(lambda v: "" if pd.isna(v) else (f"{v:.0f}" if float(v).is_integer() else f"{v:.3f}"))
    for col in table_data.select_dtypes(include=["object", "string"]):
        table_data[col] = table_data[col].map(lambda v: str(v)[:max_column_chars - 1] + "…" if pd.notna(v) and len(str(v)) > max_column_chars else v)
    table = ax.table(cellText=table_data.fillna("").astype(str).values, colLabels=table_data.columns,
                     cellLoc="left", colLoc="left", loc="center", bbox=[0, .04, 1, .82 if subtitle else .88])
    table.auto_set_font_size(False); table.set_fontsize(7.5 if len(table_data) < 24 else 6.5)
    table.scale(1, 1.2)
    for (row, _), cell in table.get_celld().items():
        cell.set_edgecolor("#d8e0e5")
        if row == 0: cell.set_facecolor("#315f7d"); cell.set_text_props(color="white", weight="bold")
        elif row % 2 == 0: cell.set_facecolor("#f1f5f7")
    add_report_footer(fig, page_number, report_label); pdf.savefig(fig, bbox_inches="tight"); plt.close(fig)


def export_excel(
    tables: dict[str, pd.DataFrame],
    path: str | Path,
    *,
    figure: object | None = None,
    figure_sheet: str | None = None,
    figure_title: str | None = None,
    figure_subtitle: str | None = None,
    figure_width: int = 1100,
    figure_height: int = 440,
    figure_dpi: int = 180,
) -> Path:
    """Write named dataframes to a styled Excel workbook (openpyxl is optional).

    Each sheet gets a bold header row with a frozen top row and an
    autofilter, plus auto-sized columns, matching the notebook's workbook
    formatting. When ``figure`` is given (for example the triple-network
    plot), it is rendered to PNG and embedded below the table on
    ``figure_sheet`` (defaulting to the last table's sheet name), with
    ``figure_title``/``figure_subtitle`` text placed above it -- reproducing
    the notebook's "Trinetwork" sheet layout.
    """
    try:
        from io import BytesIO
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ImportError("export_excel requires optional dependency 'openpyxl'") from exc

    sheet_names = {name: name[:31] for name in tables}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_names[name], index=False)

    workbook = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="315F7D")
    header_font = Font(color="FFFFFF", bold=True)
    row_counts: dict[str, int] = {}
    for name, sheet_name in sheet_names.items():
        worksheet = workbook[sheet_name]
        row_counts[name] = len(tables[name])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            width = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(width + 2, 12), 42)

    if figure is not None:
        target_name = figure_sheet or next(reversed(tables))
        if target_name not in sheet_names:
            raise KeyError(f"figure_sheet {target_name!r} is not one of the exported tables: {list(tables)}")
        worksheet = workbook[sheet_names[target_name]]
        buffer = BytesIO()
        figure.savefig(buffer, format="png", dpi=figure_dpi, bbox_inches="tight")
        buffer.seek(0)
        plot_row = row_counts[target_name] + 4
        if figure_title:
            title_cell = worksheet.cell(row=plot_row, column=1, value=figure_title)
            title_cell.font = Font(size=14, bold=True, color="315F7D")
            plot_row += 1
        if figure_subtitle:
            subtitle_cell = worksheet.cell(row=plot_row, column=1, value=figure_subtitle)
            subtitle_cell.alignment = Alignment(wrap_text=True)
            worksheet.merge_cells(start_row=plot_row, start_column=1, end_row=plot_row, end_column=6)
            plot_row += 1
        image = ExcelImage(buffer)
        image.width = figure_width
        image.height = figure_height
        worksheet.add_image(image, f"A{plot_row + 1}")

    workbook.save(path)
    return Path(path)


def export_pdf(pages: list[tuple[str, pd.DataFrame]], path: str | Path,
               subtitle: str | None = None) -> Path:
    """Export dataframe report pages to PDF using matplotlib's PdfPages."""
    try:
        from matplotlib.backends.backend_pdf import PdfPages
    except ImportError as exc:
        raise ImportError("export_pdf requires optional dependency 'matplotlib'") from exc
    with PdfPages(path) as pdf:
        for number, (title, table) in enumerate(pages, start=1):
            add_report_table_page(pdf, title, table, number, subtitle=subtitle)
    return Path(path)


def _build_triple_network_figure(
    triple_network_df: pd.DataFrame,
    triple_feature_order: Sequence[str],
    palette: dict[str, str] | None = None,
) -> object:
    """Build the triple-network boxplot figure used on the report's last page."""
    try:
        import matplotlib.pyplot as plt
        import seaborn as sns
    except ImportError as exc:
        raise ImportError(
            "_build_triple_network_figure requires optional dependencies 'matplotlib' and 'seaborn'"
        ) from exc
    palette = palette or DEFAULT_TRIPLE_PALETTE
    figure = plt.figure(figsize=(11.7, 8.3))
    ax = figure.add_subplot(111)
    sns.boxplot(
        data=triple_network_df, x="feature", y="value", hue="partition",
        order=list(triple_feature_order), hue_order=["9-network", "19-network"],
        palette=palette, showfliers=False, width=0.68, linewidth=1.1, ax=ax,
    )
    ax.set_title(
        "Triple-network connectivity by atlas resolution",
        loc="left", fontsize=16, weight="bold", color="#24465c",
    )
    ax.set_xlabel("")
    ax.set_ylabel("Mean correlation")
    ax.tick_params(axis="x", rotation=20)
    ax.legend(title="Partition", frameon=False)
    return figure


def _add_summary_page(pdf: object, summary_lines: Sequence[str], report_label: str) -> None:
    import matplotlib.pyplot as plt
    figure = plt.figure(figsize=(11.7, 8.3))
    ax = figure.add_subplot(111)
    ax.axis("off")
    figure.text(0.07, 0.91, "rs-X1 Analysis Results", fontsize=25, weight="bold", color="#24465c")
    figure.text(
        0.07, 0.86,
        "Resting-state connectivity, motion QC, outliers, and triple-network summaries",
        fontsize=11, color="#566975",
    )
    figure.text(
        0.09, 0.76, "\n".join(f"\u2022 {line}" for line in summary_lines),
        fontsize=12, va="top", linespacing=1.8,
    )
    add_report_footer(figure, 1, report_label)
    pdf.savefig(figure, bbox_inches="tight")
    plt.close(figure)


def _add_motion_association_page(
    pdf: object, motion_associations: pd.DataFrame, page_number: int,
    motion_mode: str, fd_threshold: float, report_label: str, top_n: int = 20,
) -> None:
    import matplotlib.pyplot as plt
    association_plot = motion_associations.head(top_n).sort_values("spearman_rho")
    fig, ax = plt.subplots(figsize=(11.7, 8.3))
    colors = ["#c27a62" if value < 0 else "#477c9c" for value in association_plot["spearman_rho"]]
    ax.barh(association_plot["feature"], association_plot["spearman_rho"], color=colors)
    ax.axvline(0, color="#303b42", linewidth=0.8)
    metric_label = (
        f"remaining seconds (FD={fd_threshold:g})"
        if motion_mode == "dcan_qc"
        else "num_retained_volumes"
    )
    ax.set_xlabel(f"Spearman correlation with {metric_label}")
    ax.set_title(
        "Connectivity features most associated with motion-related data loss",
        loc="left", fontsize=16, weight="bold", color="#24465c",
    )
    ax.grid(axis="x", alpha=0.2)
    add_report_footer(fig, page_number, report_label)
    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)


def _add_figure_page(pdf: object, figure: object, page_number: int, report_label: str) -> None:
    add_report_footer(figure, page_number, report_label)
    pdf.savefig(figure, bbox_inches="tight")
    import matplotlib.pyplot as plt
    plt.close(figure)


def export_analysis_results_pdf(
    path: str | Path,
    *,
    atlas_name: str,
    n_atlas_parcels: int,
    fd_threshold: float,
    n_conn_files: int,
    n_motion_files: int,
    n_motion_readable: int,
    n_master_rows: int,
    n_merged_run_rows: int,
    n_features: int,
    n_flagged_subjects: int,
    threshold_summary: pd.DataFrame,
    motion_associations: pd.DataFrame,
    outlier_feature_summary: pd.DataFrame,
    outlier_subject_report: pd.DataFrame,
    triple_summary: pd.DataFrame,
    triple_network_df: pd.DataFrame | None = None,
    triple_feature_order: Sequence[str] | None = None,
    triple_network_figure: object | None = None,
    outlier_z_threshold: float = 3.5,
    motion_mode: str = "dcan_qc",
    report_label: str = "rs-X1 analysis",
) -> Path:
    """Export the 7-page print-ready analysis results PDF.

    Pages: (1) run summary, (2) motion threshold summary table, (3) feature/
    motion association bar plot, (4) features with the most outlier flags,
    (5) outlier subject report (abbreviated feature names), (6) triple-
    network summary table, (7) triple-network boxplot. Supply either a
    ready-made ``triple_network_figure`` or ``triple_network_df`` plus
    ``triple_feature_order`` so the figure is built internally.
    """
    try:
        from matplotlib.backends.backend_pdf import PdfPages
    except ImportError as exc:
        raise ImportError(
            "export_analysis_results_pdf requires optional dependency 'matplotlib'"
        ) from exc
    if triple_network_figure is None:
        if triple_network_df is None or triple_feature_order is None:
            raise ValueError(
                "Provide triple_network_figure, or both triple_network_df and triple_feature_order"
            )
        triple_network_figure = _build_triple_network_figure(triple_network_df, triple_feature_order)

    summary_lines = [
        f"Atlas: {atlas_name} ({n_atlas_parcels} parcels)",
        "Connectivity partitions: Schaefer 9-network and 19-network",
        "ROI seeds: left, right, and bilateral hippocampus and amygdala",
        (
            f"Motion threshold used for the master features: FD={fd_threshold:g}"
            if motion_mode == "dcan_qc"
            else "Motion mode: LINC QC CSV (runs ranked by retained volumes)"
        ),
        f"Connectivity matrices found: {n_conn_files}",
        f"Motion files found: {n_motion_files} ({n_motion_readable} readable)",
        f"Matched readable subject/session/run rows: {n_master_rows}",
        f"Subject/session rows after run averaging: {n_merged_run_rows}",
        f"Connectivity features screened: {n_features}",
        f"Subjects with at least one flagged feature: {n_flagged_subjects}",
        "Outliers are review flags only; they are not automatically excluded.",
        "Run-level motion-feature associations are exploratory and do not account for repeated measures.",
    ]

    subject_report_for_pdf = outlier_subject_report.copy()
    if "flagged_features" in subject_report_for_pdf:
        subject_report_for_pdf["flagged_features"] = subject_report_for_pdf["flagged_features"].map(
            lambda names: ", ".join(str(names).split(", ")[:4])
            + (" \u2026" if len(str(names).split(", ")) > 4 else "")
        )

    with PdfPages(path) as pdf:
        _add_summary_page(pdf, summary_lines, report_label)
        add_report_table_page(
            pdf,
            "Motion Threshold Summary" if motion_mode == "dcan_qc" else "LINC Motion Summary",
            threshold_summary, 2,
            subtitle=(
                "Remaining scan duration is shown for the candidate FD thresholds; "
                f"FD={fd_threshold:g} is used for downstream features."
            ) if motion_mode == "dcan_qc" else
            "Summary of num_retained_volumes from the LINC QC CSV rows.",
            report_label=report_label,
        )
        _add_motion_association_page(
            pdf, motion_associations, 3, motion_mode, fd_threshold, report_label
        )
        add_report_table_page(
            pdf, "Features with the Most Outlier Flags", outlier_feature_summary.head(20), 4,
            subtitle=(
                f"Robust |z| > {outlier_z_threshold:g}; median/MAD scoring with an "
                "IQR fallback when MAD is zero."
            ),
            report_label=report_label,
        )
        add_report_table_page(
            pdf, "Outlier Subject Report", subject_report_for_pdf, 5,
            subtitle=(
                "Each subject is listed if any connectivity or ROI-to-network feature is "
                "flagged in any session; feature names are abbreviated for print readability."
            ),
            max_column_chars=58, report_label=report_label,
        )
        add_report_table_page(
            pdf, "Triple-Network Summary", triple_summary, 6,
            subtitle=(
                "Control (CEN), Default (DMN), and Salience (SN); 19-network values use "
                "parcel-pair-weighted macro-network means."
            ),
            report_label=report_label,
        )
        _add_figure_page(pdf, triple_network_figure, 7, report_label)
    return Path(path)
